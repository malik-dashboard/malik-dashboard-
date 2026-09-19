import streamlit as st
import pandas as pd
import io

import plotly.express as px

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter


# =========================================================
# PAGE SETTINGS
# =========================================================

st.set_page_config(
    page_title="Universal Dashboard",
    layout="wide",
    page_icon="📊"
)


# =========================================================
# HEADER
# =========================================================

st.markdown("""
<div style="
text-align:center;
background: linear-gradient(90deg, #0A1931 0%, #185ADB 100%);
padding:15px;
border-radius:10px;
color:white;
">
<h2 style="margin:0; color:white;">
📊 Universal Dashboard - AUTO HEADER
</h2>
</div>
""", unsafe_allow_html=True)


# =========================================================
# FIND BEST HEADER
# =========================================================

def find_best_header(file, sheet):

    best_row = 0
    max_score = 0

    for r in range(8):

        try:

            file.seek(0)

            temp = pd.read_excel(
                file,
                sheet_name=sheet,
                header=r
            )

            temp = temp.dropna(how="all")
            temp = temp.dropna(
                axis=1,
                how="all"
            )

            cols = len([
                c for c in temp.columns
                if "Unnamed" not in str(c)
            ])

            col_str = " ".join(
                str(c).lower()
                for c in temp.columns
            )

            score = cols

            keywords = [
                "status",
                "row labels",
                "wir",
                "system",
                "originator",
                "approved",
                "form title",
                "user ref"
            ]

            if any(
                word in col_str
                for word in keywords
            ):
                score += 10

            if score > max_score:

                max_score = score
                best_row = r

        except Exception:
            pass

    return best_row


# =========================================================
# EXCEL UPLOAD
# =========================================================

f = st.file_uploader(
    "📂 Excel Upload Karo",
    type=["xlsx", "xls"]
)

if not f:
    st.stop()


# =========================================================
# SHEET SELECT
# =========================================================

xls = pd.ExcelFile(f)

sheet = st.selectbox(
    f"Sheet Select Karo ({len(xls.sheet_names)} sheets)",
    xls.sheet_names
)


# =========================================================
# AUTO HEADER
# =========================================================

auto_header = find_best_header(
    f,
    sheet
)

st.info(
    f"🤖 Auto Detected Header Row: {auto_header} "
    f"(Aap change kar sakte ho)"
)

header_row = st.number_input(
    "Header Row Number",
    min_value=0,
    max_value=10,
    value=auto_header,
    step=1
)


# =========================================================
# READ EXCEL
# =========================================================

f.seek(0)

df = pd.read_excel(
    f,
    sheet_name=sheet,
    header=header_row
)


# =========================================================
# CLEAN DATA
# =========================================================

df = df.dropna(how="all")

# Remove Unnamed columns
valid_columns = [
    c for c in df.columns
    if not str(c).strip().lower().startswith("unnamed")
]

if valid_columns:
    df = df[valid_columns]


# =========================================================
# FALLBACK HEADER DETECTION
# =========================================================

if len(df.columns) < 2:

    for r in range(8):

        try:

            f.seek(0)

            temp = pd.read_excel(
                f,
                sheet_name=sheet,
                header=r
            )

            temp = temp.dropna(how="all")
            temp = temp.dropna(
                axis=1,
                how="all"
            )

            if len(temp.columns) >= 3:

                df = temp
                header_row = r

                break

        except Exception:
            pass


# =========================================================
# LOADED MESSAGE
# =========================================================

st.success(
    f"✅ Loaded: {len(df)} Records | "
    f"{len(df.columns)} Columns | "
    f"Header Row: {header_row}"
)

st.write(
    "Columns:",
    list(df.columns)
)


# =========================================================
# FIND STATUS COLUMN
# =========================================================

status_col = None

for c in df.columns:

    clean_name = (
        str(c)
        .strip()
        .lower()
    )

    if clean_name == "status":

        status_col = c
        break


# =========================================================
# FIND COMMON COLUMNS
# =========================================================

id_col = None
user_ref_col = None
form_title_col = None
originator_col = None


for c in df.columns:

    name = (
        str(c)
        .strip()
        .lower()
    )

    if name in ["id", "wir", "wir id"]:
        id_col = c

    if "user ref" in name:
        user_ref_col = c

    if "form title" in name:
        form_title_col = c

    if "originator" in name:
        originator_col = c


# =========================================================
# PIVOT DETECTION
# =========================================================

is_pivot = any(
    str(c).strip().lower() == "row labels"
    for c in df.columns
)


# =========================================================
# VARIABLES FOR EXCEL EXPORT
# =========================================================

excel_chart_data = []


# =========================================================
# PIVOT DASHBOARD
# =========================================================

if is_pivot:

    df_count = df[
        [
            c for c in df.columns
            if ".1" not in str(c)
        ]
    ]

    grand_cols = [
        c for c in df_count.columns
        if "grand total" in str(c).lower()
    ]

    grand_col = (
        grand_cols[0]
        if grand_cols
        else None
    )

    row_cols = [
        c for c in df_count.columns
        if "row labels" in str(c).lower()
    ]

    row_col = (
        row_cols[0]
        if row_cols
        else df_count.columns[0]
    )


    if grand_col:

        grand_idx = (
            df_count.columns.get_loc(
                grand_col
            )
        )

        status_cols = list(
            df_count.columns[
                1:grand_idx
            ]
        )

        df_data = df_count[
            (
                df_count[row_col]
                .astype(str)
                .str.lower()
                != "grand total"
            )
            &
            (
                df_count[row_col].notna()
            )
        ].copy()


        for c in status_cols + [grand_col]:

            df_data[c] = pd.to_numeric(
                df_data[c],
                errors="coerce"
            ).fillna(0)


        # -------------------------------------------------
        # METRICS
        # -------------------------------------------------

        c1, c2, c3 = st.columns(3)

        c1.metric(
            "SYSTEMS",
            len(df_data)
        )

        c2.metric(
            "TOTAL WIR",
            int(
                df_data[grand_col].sum()
            )
        )

        c3.metric(
            "STATUS TYPES",
            len(status_cols)
        )


        # -------------------------------------------------
        # TOP 10 SYSTEMS
        # -------------------------------------------------

        top10 = (
            df_data
            .sort_values(
                grand_col,
                ascending=False
            )
            .head(10)
        )


        # Save for Excel
        excel_chart_data.append({
            "type": "bar",
            "title": "Top 10 Systems - Status Wise",
            "category": row_col,
            "value": grand_col,
            "data": top10[
                [row_col, grand_col]
            ].copy()
        })


        # -------------------------------------------------
        # STATUS TOTALS
        # -------------------------------------------------

        status_totals = []

        for sc in status_cols:

            total = float(
                df_data[sc].sum()
            )

            if total > 0:

                status_totals.append({
                    "Status": str(sc),
                    "Count": int(total)
                })


        status_df = pd.DataFrame(
            status_totals
        )


        excel_chart_data.append({
            "type": "donut",
            "title": "Status Breakdown",
            "category": "Status",
            "value": "Count",
            "data": status_df.copy()
        })


        # -------------------------------------------------
        # DISPLAY
        # -------------------------------------------------

        a, b = st.columns(2)


        with a:

            fig = px.bar(
                top10,
                x=grand_col,
                y=row_col,
                orientation="h",
                text=grand_col,
                title="Top 10 Systems - Status Wise"
            )

            fig.update_layout(
                yaxis={
                    "categoryorder":
                    "total ascending"
                }
            )

            st.plotly_chart(
                fig,
                use_container_width=True
            )


        with b:

            if not status_df.empty:

                fig2 = px.pie(
                    status_df,
                    values="Count",
                    names="Status",
                    hole=0.5,
                    title="Status Breakdown"
                )

                fig2.update_traces(
                    textinfo="percent+label"
                )

                st.plotly_chart(
                    fig2,
                    use_container_width=True
                )


        final_df = df_data


    else:

        final_df = df


# =========================================================
# NORMAL EXCEL DASHBOARD
# =========================================================

else:

    # -----------------------------------------------------
    # TEXT COLUMNS
    # -----------------------------------------------------

    text_cols = df.select_dtypes(
        include=["object"]
    ).columns.tolist()


    if len(text_cols) == 0:

        text_cols = list(
            df.columns[:4]
        )


    # -----------------------------------------------------
    # METRICS
    # -----------------------------------------------------

    c1, c2, c3, c4 = st.columns(4)

    c1.metric(
        "ROWS",
        len(df)
    )

    c2.metric(
        "COLS",
        len(df.columns)
    )

    c3.metric(
        "TEXT COLS",
        len(text_cols)
    )

    c4.metric(
        "NUMERIC",
        len(
            df.select_dtypes(
                include="number"
            ).columns
        )
    )


    # =====================================================
    # IMPORTANT:
    # STATUS ALWAYS FIRST
    # =====================================================

    chart_cols = []


    # Status ko sabse pehle add karo
    if status_col is not None:

        chart_cols.append(
            status_col
        )


    # Baaki columns
    for c in text_cols:

        if c == status_col:
            continue

        if c not in chart_cols:
            chart_cols.append(c)


    # Maximum 4 charts
    chart_cols = chart_cols[:4]


    # -----------------------------------------------------
    # CHART COLUMNS
    # -----------------------------------------------------

    cols = st.columns(2)


    for i, col in enumerate(
        chart_cols
    ):

        try:

            vc = (
                df[col]
                .astype(str)
                .str.strip()
            )


            vc = vc[
                ~vc.isin([
                    "nan",
                    "None",
                    "",
                    "NaT"
                ])
            ]


            vc = (
                vc
                .value_counts()
                .head(10)
                .reset_index()
            )


            if len(vc) == 0:
                continue


            vc.columns = [
                str(col),
                "Count"
            ]


            # -------------------------------------------------
            # SAVE DATA FOR EXCEL
            # -------------------------------------------------

            if col == status_col:

                excel_chart_data.append({
                    "type": "donut",
                    "title": "Status Breakdown",
                    "category": str(col),
                    "value": "Count",
                    "data": vc.copy()
                })

            else:

                excel_chart_data.append({
                    "type": "bar",
                    "title": f"{col} - Top 10",
                    "category": str(col),
                    "value": "Count",
                    "data": vc.copy()
                })


            # -------------------------------------------------
            # STREAMLIT CHART
            # -------------------------------------------------

            with cols[i % 2]:

                # =============================================
                # STATUS = DONUT
                # =============================================

                if col == status_col:

                    fig = px.pie(
                        vc,
                        values="Count",
                        names=str(col),
                        hole=0.5,
                        title="Status Breakdown"
                    )

                    fig.update_traces(
                        textinfo="percent+label"
                    )


                # =============================================
                # EVERYTHING ELSE = BAR
                # =============================================

                else:

                    fig = px.bar(
                        vc,
                        x="Count",
                        y=str(col),
                        orientation="h",
                        text="Count",
                        title=f"{col} - Top 10"
                    )

                    fig.update_layout(
                        yaxis={
                            "categoryorder":
                            "total ascending"
                        }
                    )


                st.plotly_chart(
                    fig,
                    use_container_width=True
                )


        except Exception:
            pass


    final_df = df


# =========================================================
# DATA PREVIEW
# =========================================================

st.markdown(
    "### 📋 Data Preview"
)

st.dataframe(
    final_df.head(50),
    use_container_width=True
)


# =========================================================
# EXCEL EXPORT FUNCTION
# =========================================================

def get_excel():

    # -----------------------------------------------------
    # CREATE INITIAL EXCEL
    # -----------------------------------------------------

    out = io.BytesIO()


    with pd.ExcelWriter(
        out,
        engine="openpyxl"
    ) as writer:

        final_df.to_excel(
            writer,
            index=False,
            sheet_name="Data"
        )


    out.seek(0)


    # -----------------------------------------------------
    # LOAD WORKBOOK
    # -----------------------------------------------------

    wb = load_workbook(
        out
    )


    # Remove if already exists

    if "Dashboard" in wb.sheetnames:
        del wb["Dashboard"]

    if "Summary" in wb.sheetnames:
        del wb["Summary"]


    # Create sheets

    ws_dash = wb.create_sheet(
        "Dashboard",
        0
    )

    ws_data = wb["Data"]

    ws_summary = wb.create_sheet(
        "Summary"
    )


    # =====================================================
    # DASHBOARD TITLE
    # =====================================================

    ws_dash["A1"] = (
        "📊 Universal Dashboard"
    )

    ws_dash["A1"].font = Font(
        size=20,
        bold=True,
        color="FFFFFF"
    )

    ws_dash["A1"].fill = PatternFill(
        "solid",
        fgColor="185ADB"
    )

    ws_dash["A1"].alignment = Alignment(
        horizontal="center"
    )

    ws_dash.merge_cells(
        "A1:L2"
    )


    # =====================================================
    # DASHBOARD METRICS
    # =====================================================

    ws_dash["A4"] = "Total Records"
    ws_dash["B4"] = len(final_df)

    ws_dash["D4"] = "Total Columns"
    ws_dash["E4"] = len(final_df.columns)

    ws_dash["G4"] = "Status Column"
    ws_dash["H4"] = (
        str(status_col)
        if status_col
        else "Not Found"
    )


    for cell in [
        "A4",
        "B4",
        "D4",
        "E4",
        "G4",
        "H4"
    ]:

        ws_dash[cell].font = Font(
            bold=True
        )


    # =====================================================
    # WRITE CHART DATA + CREATE EXCEL CHARTS
    # =====================================================

    summary_col = 1

    chart_positions = [
        "A7",
        "J7",
        "A25",
        "J25"
    ]


    chart_number = 0


    for item in excel_chart_data:

        data_df = item["data"]


        if data_df is None:
            continue

        if data_df.empty:
            continue


        title = item["title"]

        category = item["category"]

        value = item["value"]

        chart_type = item["type"]


        # -------------------------------------------------
        # SUMMARY DATA
        # -------------------------------------------------

        start_col = summary_col

        ws_summary.cell(
            1,
            start_col,
            category
        )

        ws_summary.cell(
            1,
            start_col + 1,
            value
        )


        for r, (_, row) in enumerate(
            data_df.iterrows(),
            start=2
        ):

            ws_summary.cell(
                r,
                start_col,
                str(row.iloc[0])
            )

            ws_summary.cell(
                r,
                start_col + 1,
                float(row.iloc[1])
            )


        max_row = len(data_df) + 1


        # -------------------------------------------------
        # CREATE DONUT
        # -------------------------------------------------

        if chart_type == "donut":

            chart = DoughnutChart()

            chart.title = title

            chart.holeSize = 50

            chart.height = 8
            chart.width = 13


            data = Reference(
                ws_summary,
                min_col=start_col + 1,
                min_row=1,
                max_row=max_row
            )

            labels = Reference(
                ws_summary,
                min_col=start_col,
                min_row=2,
                max_row=max_row
            )


            chart.add_data(
                data,
                titles_from_data=True
            )

            chart.set_categories(
                labels
            )


            chart.dataLabels = DataLabelList()

            chart.dataLabels.showPercent = True

            chart.dataLabels.showLeaderLines = True


        # -------------------------------------------------
        # CREATE BAR
        # -------------------------------------------------

        else:

            chart = BarChart()

            chart.type = "bar"

            chart.style = 10

            chart.title = title

            chart.x_axis.title = "Count"

            chart.y_axis.title = category

            chart.height = 8
            chart.width = 14


            data = Reference(
                ws_summary,
                min_col=start_col + 1,
                min_row=1,
                max_row=max_row
            )

            labels = Reference(
                ws_summary,
                min_col=start_col,
                min_row=2,
                max_row=max_row
            )


            chart.add_data(
                data,
                titles_from_data=True
            )

            chart.set_categories(
                labels
            )


        # -------------------------------------------------
        # ADD CHART TO DASHBOARD
        # -------------------------------------------------

        if chart_number < len(
            chart_positions
        ):

            ws_dash.add_chart(
                chart,
                chart_positions[
                    chart_number
                ]
            )


        chart_number += 1

        summary_col += 3


    # =====================================================
    # FORMAT DATA SHEET
    # =====================================================

    for ws in [
        ws_data,
        ws_summary
    ]:

        ws.freeze_panes = "A2"


        # Header
        for cell in ws[1]:

            cell.font = Font(
                bold=True,
                color="FFFFFF"
            )

            cell.fill = PatternFill(
                "solid",
                fgColor="185ADB"
            )

            cell.alignment = Alignment(
                horizontal="center"
            )


        # Auto width
        for column_cells in ws.columns:

            max_length = 0

            column_letter = get_column_letter(
                column_cells[0].column
            )


            for cell in column_cells:

                try:

                    length = len(
                        str(cell.value)
                    )

                    if length > max_length:
                        max_length = length

                except Exception:
                    pass


            ws.column_dimensions[
                column_letter
            ].width = min(
                max_length + 2,
                40
            )


    # =====================================================
    # FORMAT DASHBOARD
    # =====================================================

    for col in range(1, 14):

        ws_dash.column_dimensions[
            get_column_letter(col)
        ].width = 14


    ws_dash.freeze_panes = "A6"


    # =====================================================
    # SAVE FINAL EXCEL
    # =====================================================

    final_output = io.BytesIO()

    wb.save(
        final_output
    )

    final_output.seek(0)

    return final_output.getvalue()


# =========================================================
# DOWNLOAD BUTTON
# =========================================================

st.markdown("---")

st.download_button(
    label=(
        f"📥 DOWNLOAD EXCEL + CHARTS "
        f"- {len(final_df)} Records"
    ),

    data=get_excel(),

    file_name=(
        f"Report_{len(final_df)}_Dashboard.xlsx"
    ),

    mime=(
        "application/vnd.openxmlformats-officedocument."
        "spreadsheetml.sheet"
    ),

    use_container_width=True,

    type="primary"
)
